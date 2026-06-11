"""Export a formatted Excel variance report with openpyxl.

Showing you can automate the deliverable a manager actually wants -- a clean,
color-coded variance workbook -- is a meaningful differentiator for an analyst
role where Excel is the lingua franca.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .variance import variance_summary

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FAV_FILL = PatternFill("solid", fgColor="E2EFDA")
UNFAV_FILL = PatternFill("solid", fgColor="FCE4D6")
TOTAL_FONT = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def export_variance_report(report: pd.DataFrame, out_path: str | Path) -> Path:
    out_path = Path(out_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Variance Report"

    month = report["month"].iloc[0] if not report.empty else ""
    ws["A1"] = f"Budget vs. Actual Variance — {month}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:F1")

    headers = ["Category", "Budget", "Actual", "Variance", "Variance %", "Status"]
    for col, name in enumerate(headers, start=1):
        c = ws.cell(row=3, column=col, value=name)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER

    row = 4
    for _, r in report.iterrows():
        fill = FAV_FILL if r["status"] == "Favorable" else UNFAV_FILL
        values = [
            r["category"], r["budget"], r["actual"], r["variance"],
            r["variance_pct"], r["status"],
        ]
        for col, val in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=val)
            c.fill = fill
            c.border = BORDER
            if col in (2, 3, 4):
                c.number_format = "#,##0.00"
            if col == 5:
                c.number_format = "0.0%"
        row += 1

    summary = variance_summary(report)
    if summary:
        ws.cell(row=row, column=1, value="TOTAL").font = TOTAL_FONT
        for col, key in zip((2, 3, 4), ("total_budget", "total_actual", "total_variance")):
            c = ws.cell(row=row, column=col, value=summary[key])
            c.font = TOTAL_FONT
            c.number_format = "#,##0.00"
        pct = ws.cell(row=row, column=5, value=summary["total_variance_pct"])
        pct.font = TOTAL_FONT
        pct.number_format = "0.0%"
        ws.cell(row=row, column=6, value=summary["status"]).font = TOTAL_FONT

    for col in range(1, 7):
        ws.column_dimensions[get_column_letter(col)].width = 16

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
